"""
Knowledge base: one library of files per firm (workspace), searched by the agents on calls.

An admin uploads a file (PDF, Word, Excel, CSV, text, Markdown or an image). It is read ONCE, here, into
plain text, split into short overlapping passages and indexed; calls only search those passages, so the
call worker needs none of the readers below. Images, and PDF pages that hold no text (scans), are read
by Gemini; without a GEMINI_API_KEY those files fail with a plain reason instead of sitting empty.

Storage follows agent/case_manager.py: with DATABASE_URL the file records and passages live in
``app_documents`` (collections ``knowledge_files`` and ``knowledge_chunks``) and the original bytes in
``knowledge_blobs``; without it everything is JSON and files under config/knowledge/. Another process
(the worker) picks up changes within REFRESH_SECONDS.
"""

from __future__ import annotations

import io
import json
import logging
import math
import os
import re
import threading
import time
import uuid
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

log = logging.getLogger("knowledge")

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOCAL_DIR = os.path.join(ROOT_DIR, "config", "knowledge")

FILES = "knowledge_files"
CHUNKS = "knowledge_chunks"

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_FIRM_BYTES = 200 * 1024 * 1024
MAX_FILES_PER_FIRM = 300
MAX_TEXT_CHARS = 400_000          # per file, after reading
MAX_SCANNED_PAGES = 20            # PDF pages without text sent to Gemini, per file
MAX_UNZIPPED_BYTES = 150 * 1024 * 1024   # .docx/.xlsx are zip files; refuse zip bombs
MAX_SHEET_ROWS = 5000
CHUNK_CHARS = 900
CHUNK_OVERLAP = 150
REFRESH_SECONDS = 5

KINDS = {
    "pdf": "PDF", "docx": "Word document", "xlsx": "Excel workbook", "csv": "CSV",
    "markdown": "Markdown", "text": "Text", "image": "Image",
}
STATUSES = ("processing", "ready", "failed")


class KnowledgeError(ValueError):
    """A refusal with a message fit to show the person who uploaded the file."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


@dataclass
class KnowledgeFile:
    file_id: str
    workspace_id: str
    name: str
    kind: str
    mime: str
    size_bytes: int
    sha256: str
    status: str = "processing"
    error: str = ""
    notes: str = ""                 # e.g. "2 scanned pages could not be read"
    engine: str = ""                # what read it: pdf-text, pdf-text+gemini, docx, xlsx, gemini, text
    pages: int = 0
    chars: int = 0
    chunk_count: int = 0
    uploaded_by: str = ""
    created_at: float = field(default_factory=time.time)
    updated_at: float = field(default_factory=time.time)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

    def public(self) -> Dict[str, Any]:
        d = self.to_dict()
        d["kind_label"] = KINDS.get(self.kind, self.kind)
        return d


# ── File type: decided from the bytes, never from the name alone ─────────────────────────────────────
def sniff(data: bytes, name: str) -> Tuple[str, str]:
    """(kind, mime) for an upload, or KnowledgeError saying why it isn't accepted."""
    ext = os.path.splitext(name or "")[1].lower()
    if not data:
        raise KnowledgeError("The file is empty.")
    if data.startswith(b"%PDF-"):
        return "pdf", "application/pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image", "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image", "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image", "image/webp"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise KnowledgeError("Old Word/Excel files (.doc, .xls) aren't supported. Save it as .docx or .xlsx and upload that.")
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
                unzipped = sum(i.file_size for i in z.infolist())
        except zipfile.BadZipFile:
            raise KnowledgeError("The file looks damaged and can't be opened.")
        if unzipped > MAX_UNZIPPED_BYTES:
            raise KnowledgeError("The file expands to more than 150 MB when opened, so it isn't accepted.")
        if "word/document.xml" in names:
            return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if "xl/workbook.xml" in names:
            return "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        raise KnowledgeError("Zip files and this kind of Office file aren't supported. Upload PDF, Word, Excel, CSV, text or an image.")
    # Plain text: must decode and must not look binary.
    head = data[:65536]
    if b"\x00" in head:
        raise KnowledgeError("This file type isn't supported. Upload PDF, Word (.docx), Excel (.xlsx), CSV, text, Markdown, or a PNG/JPG/WebP image.")
    try:
        _decode_text(data)
    except UnicodeDecodeError:
        raise KnowledgeError("The text file isn't UTF-8 encoded. Save it as UTF-8 and upload it again.")
    if ext == ".csv":
        return "csv", "text/csv"
    if ext in (".md", ".markdown"):
        return "markdown", "text/markdown"
    return "text", "text/plain"


def _decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    return data.decode("utf-8")


def clean_name(name: str) -> str:
    base = os.path.basename((name or "").replace("\\", "/")).strip()
    base = re.sub(r"[\x00-\x1f\x7f]", "", base)
    return (base or "Untitled")[:120]


# ── Reading files into text ───────────────────────────────────────────────────────────────────────────
IMAGE_SYSTEM = (
    "You read images uploaded to a law firm's knowledge base so a phone receptionist can answer callers from them. "
    "Return JSON: {\"text\": every piece of readable text in the image, verbatim, in reading order (tables as rows "
    "with ' | ' between cells), or \"\" if there is none; \"description\": one or two plain sentences on what the "
    "image shows}. Never add facts that are not in the image."
)


def _gemini_read_image(mime: str, data: bytes) -> str:
    """Text + description of an image via Gemini. Raises RuntimeError when it can't."""
    from agent.schema_extractor import gemini_json
    mime, data = _shrink_image(mime, data)
    # Reads run in the background, so allow for the free tier's occasional slow answer: one retry.
    for attempt in (1, 2):
        try:
            out = gemini_json(IMAGE_SYSTEM, "Read this image.", max_tokens=4000, timeout=90.0, attachments=[(mime, data)])
            break
        except RuntimeError as e:
            if attempt == 2 or not any(w in str(e) for w in ("timed out", "HTTP 429", "HTTP 500", "HTTP 503")):
                raise
            time.sleep(3)
    if not isinstance(out, dict):
        raise RuntimeError("Gemini returned an unexpected answer")
    text = str(out.get("text") or "").strip()
    desc = str(out.get("description") or "").strip()
    return "\n\n".join(p for p in (desc and f"Image: {desc}", text) if p)


# Swappable for tests (a stand-in that needs no key or network).
IMAGE_READER: Callable[[str, bytes], str] = _gemini_read_image


def _shrink_image(mime: str, data: bytes) -> Tuple[str, bytes]:
    """Keeps requests small: images over 2000px or 4 MB are re-encoded as JPEG."""
    if len(data) <= 4 * 1024 * 1024:
        try:
            from PIL import Image
            with Image.open(io.BytesIO(data)) as im:
                if max(im.size) <= 2000:
                    return mime, data
        except Exception:
            return mime, data
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as im:
            im = im.convert("RGB")
            im.thumbnail((2000, 2000))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=85)
            return "image/jpeg", buf.getvalue()
    except Exception:
        return mime, data


def _image_reader_ready() -> bool:
    if IMAGE_READER is not _gemini_read_image:
        return True
    from agent.schema_extractor import _gemini_key
    return bool(_gemini_key())


def extract(kind: str, mime: str, data: bytes) -> Dict[str, Any]:
    """Reads a file. Returns {"segments": [(text, page or 0)], "pages", "engine", "notes"}.
    Raises KnowledgeError when nothing usable comes out."""
    if kind == "pdf":
        return _read_pdf(data)
    if kind == "docx":
        return {"segments": [(_read_docx(data), 0)], "pages": 0, "engine": "docx", "notes": ""}
    if kind == "xlsx":
        return {"segments": [(_read_xlsx(data), 0)], "pages": 0, "engine": "xlsx", "notes": ""}
    if kind in ("csv", "markdown", "text"):
        return {"segments": [(_decode_text(data), 0)], "pages": 0, "engine": "text", "notes": ""}
    if kind == "image":
        if not _image_reader_ready():
            raise KnowledgeError("Images are read by Gemini, and no Gemini API key is set up. Add one under API Keys & Providers, then choose Read again.")
        try:
            text = IMAGE_READER(mime, data)
        except Exception as e:
            raise KnowledgeError(f"The image couldn't be read: {e}")
        return {"segments": [(text, 0)], "pages": 0, "engine": "gemini", "notes": ""}
    raise KnowledgeError(f"Unknown file kind '{kind}'.")


def _read_pdf(data: bytes) -> Dict[str, Any]:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # older PyMuPDF
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise KnowledgeError(f"The PDF couldn't be opened: {e}")
    if doc.needs_pass:
        raise KnowledgeError("The PDF is password-protected. Upload a copy without a password.")
    segments: List[Tuple[str, int]] = []
    scanned, read_scans, unread = 0, 0, 0
    can_read_images = _image_reader_ready()
    try:
        for i, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if len(text) >= 25:
                segments.append((text, i))
                continue
            scanned += 1
            if not can_read_images or read_scans >= MAX_SCANNED_PAGES:
                unread += 1
                continue
            try:
                png = page.get_pixmap(dpi=150).tobytes("png")
                page_text = IMAGE_READER("image/png", png).strip()
                read_scans += 1
                if page_text:
                    segments.append((page_text, i))
            except Exception as e:
                log.info("Scanned page %d not read: %s", i, e)
                unread += 1
        pages = doc.page_count
    finally:
        doc.close()
    notes = []
    if unread:
        why = "no Gemini API key is set up" if not can_read_images else f"only {MAX_SCANNED_PAGES} scanned pages are read per file"
        notes.append(f"{unread} scanned page{'s' if unread != 1 else ''} could not be read ({why})")
    if not segments:
        if scanned and not can_read_images:
            raise KnowledgeError("This PDF is scanned images with no text, and no Gemini API key is set up to read them.")
        raise KnowledgeError("No text could be found in this PDF.")
    engine = "pdf-text+gemini" if read_scans else "pdf-text"
    return {"segments": segments, "pages": pages, "engine": engine, "notes": "; ".join(notes)}


def _read_docx(data: bytes) -> str:
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise KnowledgeError(f"The Word document couldn't be opened: {e}")
    parts: List[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
        parts.append("")
    return "\n\n".join(parts).strip()


def _read_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise KnowledgeError(f"The Excel workbook couldn't be opened: {e}")
    parts: List[str] = []
    try:
        for ws in wb.worksheets:
            rows: List[str] = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    rows.append(f"(rows after {MAX_SHEET_ROWS} not read)")
                    break
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                parts.append(f"Sheet: {ws.title}\n" + "\n".join(rows))
    finally:
        wb.close()
    return "\n\n".join(parts)


# ── Passages ──────────────────────────────────────────────────────────────────────────────────────────
def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def split_passages(segments: List[Tuple[str, int]]) -> List[Dict[str, Any]]:
    """Short overlapping passages, never across a PDF page. Each: {"i", "text", "page"}."""
    out: List[Dict[str, Any]] = []
    for raw, page in segments:
        text = _normalize(raw)
        if not text:
            continue
        # Paragraphs, then sentences, then words: pieces no longer than CHUNK_CHARS.
        pieces: List[str] = []
        for para in text.split("\n\n"):
            if len(para) <= CHUNK_CHARS:
                pieces.append(para)
                continue
            for sent in re.split(r"(?<=[.!?])\s+|\n", para):
                while len(sent) > CHUNK_CHARS:
                    cut = sent.rfind(" ", 0, CHUNK_CHARS)
                    cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
                    pieces.append(sent[:cut])
                    sent = sent[cut:].lstrip()
                if sent:
                    pieces.append(sent)
        buf = ""
        for piece in pieces:
            if buf and len(buf) + 2 + len(piece) > CHUNK_CHARS:
                out.append({"i": len(out), "text": buf, "page": page})
                tail = buf[-CHUNK_OVERLAP:]
                space = tail.find(" ")
                buf = (tail[space + 1:] if space >= 0 else tail) + "\n" + piece
            else:
                buf = f"{buf}\n\n{piece}" if buf else piece
        if buf:
            out.append({"i": len(out), "text": buf, "page": page})
    return out


# ── Search (BM25 over the firm's passages) ────────────────────────────────────────────────────────────
STOPWORDS = set("""a an and are as at be but by do does for from had has have how i if in into is it its me my
of on or our so than that the their them then there these they this to was we were what which who
with you your yours about any could would should please tell know""".split())

# Callers word things differently from the documents ("how much" vs "fee schedule"). A query word also
# matches the others in its group, at half weight. Kept to what a firm's receptionist is asked.
SYNONYMS = [
    "cost price fee charge rate dollar pay paid payment expensive cheap much afford billing bill retainer contingency",
    "hour open close closed time weekend saturday sunday holiday",
    "where address located location office directions find street",
    "park parking garage lot",
    "phone number call telephone fax",
    "email mail contact reach",
    "free consult consultation evaluation",
    "lawyer attorney counsel partner associate",
    "spanish espanol bilingual language",
    "case matter claim lawsuit",
]


def stem(w: str) -> str:
    """Light suffix stripping so parking/park, charged/charge and offices/office meet."""
    if w.isdigit() or len(w) <= 3:
        return w
    if w.endswith("ies") and len(w) > 4:
        return w[:-3] + "y"
    for suf in ("ing", "ed", "es", "s"):
        if w.endswith(suf) and len(w) - len(suf) >= 3 and not (suf == "s" and w.endswith("ss")):
            w = w[: -len(suf)]
            break
    if w.endswith("e") and len(w) > 3:
        w = w[:-1]
    return w


def tokens(text: str) -> List[str]:
    return [stem(w) for w in re.findall(r"[a-z0-9]+", text.lower())
            if w not in STOPWORDS and (len(w) >= 2 or w.isdigit())]


_SYN: Dict[str, List[str]] = {}
for _group in SYNONYMS:
    _words = [stem(w) for w in _group.split()]
    for _w in _words:
        _SYN.setdefault(_w, [])
        _SYN[_w] += [x for x in _words if x != _w and x not in _SYN[_w]]


class _Index:
    def __init__(self, passages: List[Dict[str, Any]]):
        self.passages = passages
        self.tf = [Counter(tokens(p["text"] + " " + p.get("file_name", ""))) for p in passages]
        self.lengths = [sum(t.values()) for t in self.tf]
        self.avg = (sum(self.lengths) / len(self.lengths)) if self.lengths else 0.0
        df: Counter = Counter()
        for t in self.tf:
            df.update(t.keys())
        n = len(passages)
        self.idf = {w: math.log(1 + (n - c + 0.5) / (c + 0.5)) for w, c in df.items()}

    def search(self, query: str, limit: int) -> List[Tuple[float, Dict[str, Any]]]:
        weights: Dict[str, float] = {}
        for w in tokens(query):
            weights[w] = 1.0
        for w in list(weights):
            for syn in _SYN.get(w, []):
                weights.setdefault(syn, 0.5)
        if not weights or not self.passages:
            return []
        k1, b = 1.5, 0.75
        scored = []
        for tf, length, p in zip(self.tf, self.lengths, self.passages):
            s = 0.0
            for w, weight in weights.items():
                f = tf.get(w)
                if f:
                    s += weight * self.idf[w] * f * (k1 + 1) / (f + k1 * (1 - b + b * length / (self.avg or 1)))
            if s > 0:
                scored.append((s, p))
        scored.sort(key=lambda x: -x[0])
        return scored[:limit]


# ── Manager ───────────────────────────────────────────────────────────────────────────────────────────
class KnowledgeManager:
    def __init__(self, local_dir: str = LOCAL_DIR):
        self.local_dir = local_dir
        self._lock = threading.RLock()
        self._files: Dict[str, KnowledgeFile] = {}
        self._chunks: Dict[str, Dict[str, Any]] = {}     # file_id -> {"text", "chunks"}
        self._indexes: Dict[str, _Index] = {}             # workspace_id -> index
        self._loaded_at = 0.0
        self._checked_at = 0.0
        self._stamps: Optional[Dict[str, float]] = None
        self._local_mtime = 0.0
        self._load()
        if os.getenv("DATABASE_URL"):
            from agent import storage
            self._stamps = storage.collection_stamps((FILES, CHUNKS))
            self._checked_at = time.time()

    # storage ------------------------------------------------------------------------------------------
    def _db(self) -> bool:
        if not os.getenv("DATABASE_URL"):
            return False
        from agent import storage
        return storage.available()

    @property
    def _index_path(self) -> str:
        return os.path.join(self.local_dir, "files.json")

    def _load(self) -> None:
        files: List[Dict[str, Any]] = []
        chunks: List[Dict[str, Any]] = []
        if os.getenv("DATABASE_URL"):
            from agent import storage
            if not storage.available():
                return       # configured but down: keep what we have rather than wipe it
            files = storage.load_collection(FILES)
            chunks = storage.load_collection(CHUNKS)
        elif os.path.isfile(self._index_path):
            try:
                self._local_mtime = os.path.getmtime(self._index_path)
                with open(self._index_path, "r", encoding="utf-8") as f:
                    files = json.load(f).get("files") or []
                for d in files:
                    p = os.path.join(self.local_dir, "chunks", d["file_id"] + ".json")
                    if os.path.isfile(p):
                        with open(p, "r", encoding="utf-8") as f:
                            chunks.append(json.load(f))
            except Exception as e:
                log.warning("Failed to load the knowledge base from %s: %s", self.local_dir, e)
                return
        known = set(KnowledgeFile.__dataclass_fields__)
        loaded: Dict[str, KnowledgeFile] = {}
        for d in files:
            try:
                kf = KnowledgeFile(**{k: v for k, v in d.items() if k in known})
                loaded[kf.file_id] = kf
            except TypeError as e:
                log.warning("Skipping malformed knowledge file record: %s", e)
        with self._lock:
            self._files = loaded
            self._chunks = {c["file_id"]: c for c in chunks if c.get("file_id") in loaded}
            self._indexes = {}
            self._loaded_at = time.time()

    def _refresh(self) -> None:
        if os.getenv("DATABASE_URL"):
            # Cheap check (two stamps) at most every REFRESH_SECONDS; a full reload only when they moved.
            if time.time() - self._checked_at < REFRESH_SECONDS:
                return
            self._checked_at = time.time()
            from agent import storage
            stamps = storage.collection_stamps((FILES, CHUNKS))
            if stamps is not None and stamps != self._stamps:
                self._load()
                self._stamps = stamps
        elif os.path.isfile(self._index_path) and os.path.getmtime(self._index_path) != self._local_mtime:
            self._load()

    def _write_local_index(self) -> None:
        os.makedirs(self.local_dir, exist_ok=True)
        tmp = self._index_path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump({"files": [k.to_dict() for k in self._files.values()], "updated_at": time.time()}, f, indent=2)
        os.replace(tmp, self._index_path)
        self._local_mtime = os.path.getmtime(self._index_path)

    def _save_file(self, kf: KnowledgeFile) -> None:
        kf.updated_at = time.time()
        if os.getenv("DATABASE_URL"):
            from agent import storage
            if not storage.save_document(FILES, kf.file_id, kf.to_dict(), kf.updated_at):
                raise KnowledgeError("The file couldn't be saved: the database is unavailable.", 503)
            return
        self._write_local_index()

    def _save_chunks(self, file_id: str, doc: Dict[str, Any]) -> None:
        if os.getenv("DATABASE_URL"):
            from agent import storage
            if not storage.save_document(CHUNKS, file_id, doc):
                raise KnowledgeError("The file's text couldn't be saved: the database is unavailable.", 503)
            return
        d = os.path.join(self.local_dir, "chunks")
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, file_id + ".json"), "w", encoding="utf-8") as f:
            json.dump(doc, f)

    def _save_blob(self, kf: KnowledgeFile, data: bytes) -> None:
        if os.getenv("DATABASE_URL"):
            from agent import storage
            if not storage.save_blob(kf.file_id, kf.workspace_id, data):
                raise KnowledgeError("The file couldn't be saved: the database is unavailable.", 503)
            return
        d = os.path.join(self.local_dir, "blobs")
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, kf.file_id + ".bin"), "wb") as f:
            f.write(data)

    def load_blob(self, file_id: str) -> Optional[bytes]:
        if os.getenv("DATABASE_URL"):
            from agent import storage
            return storage.load_blob(file_id)
        p = os.path.join(self.local_dir, "blobs", file_id + ".bin")
        if not os.path.isfile(p):
            return None
        with open(p, "rb") as f:
            return f.read()

    # reads --------------------------------------------------------------------------------------------
    def get(self, file_id: str) -> Optional[KnowledgeFile]:
        self._refresh()
        return self._files.get(file_id)

    def list_files(self, workspace_id: str) -> List[KnowledgeFile]:
        self._refresh()
        return sorted((k for k in self._files.values() if k.workspace_id == workspace_id), key=lambda k: -k.created_at)

    def usage(self, workspace_id: str) -> Dict[str, Any]:
        files = self.list_files(workspace_id)
        return {"files": len(files), "bytes": sum(k.size_bytes for k in files),
                "max_file_bytes": MAX_FILE_BYTES, "max_firm_bytes": MAX_FIRM_BYTES, "max_files": MAX_FILES_PER_FIRM,
                "ready": sum(1 for k in files if k.status == "ready"),
                "processing": sum(1 for k in files if k.status == "processing"),
                "failed": sum(1 for k in files if k.status == "failed")}

    def text(self, file_id: str) -> str:
        self._refresh()
        return (self._chunks.get(file_id) or {}).get("text", "")

    def passages(self, file_id: str) -> List[Dict[str, Any]]:
        self._refresh()
        return list((self._chunks.get(file_id) or {}).get("chunks", []))

    # writes -------------------------------------------------------------------------------------------
    def upload(self, workspace_id: str, name: str, data: bytes, uploaded_by: str = "", wait: bool = False) -> KnowledgeFile:
        """Checks and stores an upload, then reads it (in the background unless wait=True)."""
        if not workspace_id:
            raise KnowledgeError("No firm selected for this upload.")
        if len(data) > MAX_FILE_BYTES:
            raise KnowledgeError(f"The file is {len(data) / 1048576:.1f} MB; the limit is {MAX_FILE_BYTES // 1048576} MB per file.", 413)
        name = clean_name(name)
        kind, mime = sniff(data, name)
        import hashlib
        digest = hashlib.sha256(data).hexdigest()
        with self._lock:
            self._refresh()
            existing = [k for k in self._files.values() if k.workspace_id == workspace_id]
            dup = next((k for k in existing if k.sha256 == digest), None)
            if dup:
                raise KnowledgeError(f"This file is already in the knowledge base as \"{dup.name}\".", 409)
            if len(existing) >= MAX_FILES_PER_FIRM:
                raise KnowledgeError(f"The knowledge base already holds {MAX_FILES_PER_FIRM} files. Delete some before adding more.", 413)
            used = sum(k.size_bytes for k in existing)
            if used + len(data) > MAX_FIRM_BYTES:
                left = max(0, MAX_FIRM_BYTES - used) / 1048576
                raise KnowledgeError(f"Not enough space: {left:.1f} MB left of the firm's {MAX_FIRM_BYTES // 1048576} MB.", 413)
            kf = KnowledgeFile(file_id="kf-" + uuid.uuid4().hex[:12], workspace_id=workspace_id, name=name, kind=kind,
                               mime=mime, size_bytes=len(data), sha256=digest, uploaded_by=uploaded_by)
            self._save_blob(kf, data)
            self._files[kf.file_id] = kf
            try:
                self._save_file(kf)
            except Exception:
                self._files.pop(kf.file_id, None)
                raise
        if wait:
            self.process(kf.file_id, data)
        else:
            threading.Thread(target=self.process, args=(kf.file_id, data), daemon=True, name=f"knowledge-{kf.file_id}").start()
        return self._files.get(kf.file_id, kf)

    def process(self, file_id: str, data: Optional[bytes] = None) -> Optional[KnowledgeFile]:
        """Reads a stored file into passages. Never raises; the record ends ready or failed."""
        kf = self._files.get(file_id)
        if not kf:
            return None
        try:
            if data is None:
                data = self.load_blob(file_id)
                if data is None:
                    raise KnowledgeError("The original file is missing, so it can't be read again. Delete it and upload it again.")
            result = extract(kf.kind, kf.mime, data)
            passages = split_passages(result["segments"])
            text = "\n\n".join(_normalize(t) for t, _ in result["segments"] if t.strip())
            if not passages:
                raise KnowledgeError("No text could be found in this file.")
            notes = result.get("notes", "")
            if len(text) > MAX_TEXT_CHARS:
                text = text[:MAX_TEXT_CHARS]
                kept, total = [], 0
                for p in passages:
                    total += len(p["text"])
                    if total > MAX_TEXT_CHARS:
                        break
                    kept.append(p)
                passages = kept
                notes = "; ".join(x for x in (notes, f"only the first {MAX_TEXT_CHARS:,} characters were kept") if x)
            for p in passages:
                p["file_id"], p["file_name"] = kf.file_id, kf.name
            with self._lock:
                if file_id not in self._files:        # deleted while it was being read
                    return None
                self._save_chunks(file_id, {"file_id": file_id, "workspace_id": kf.workspace_id, "text": text, "chunks": passages})
                self._chunks[file_id] = {"file_id": file_id, "workspace_id": kf.workspace_id, "text": text, "chunks": passages}
                kf.status, kf.error, kf.notes, kf.engine = "ready", "", notes, result.get("engine", "")
                kf.pages, kf.chars, kf.chunk_count = int(result.get("pages") or 0), len(text), len(passages)
                self._save_file(kf)
                self._indexes.pop(kf.workspace_id, None)
        except Exception as e:
            msg = str(e) if isinstance(e, KnowledgeError) else f"Reading the file failed: {e}"
            log.warning("Knowledge file %s (%s) failed: %s", file_id, kf.name, msg)
            with self._lock:
                if file_id in self._files:
                    kf.status, kf.error, kf.chars, kf.chunk_count = "failed", msg, 0, 0
                    self._chunks.pop(file_id, None)
                    self._indexes.pop(kf.workspace_id, None)
                    try:
                        self._save_file(kf)
                    except Exception as save_err:
                        log.warning("Could not record the failure for %s: %s", file_id, save_err)
        return self._files.get(file_id)

    def reprocess(self, file_id: str, wait: bool = False) -> KnowledgeFile:
        kf = self._files.get(file_id)
        if not kf:
            raise KnowledgeError("File not found.", 404)
        with self._lock:
            kf.status, kf.error = "processing", ""
            self._save_file(kf)
        if wait:
            self.process(file_id)
        else:
            threading.Thread(target=self.process, args=(file_id,), daemon=True).start()
        return kf

    def resume_unfinished(self) -> int:
        """Re-reads files left 'processing' by a restart. Returns how many were started."""
        stuck = [k.file_id for k in self._files.values() if k.status == "processing"]
        for fid in stuck:
            threading.Thread(target=self.process, args=(fid,), daemon=True).start()
        return len(stuck)

    def delete(self, file_id: str) -> bool:
        with self._lock:
            kf = self._files.pop(file_id, None)
            if not kf:
                return False
            self._chunks.pop(file_id, None)
            self._indexes.pop(kf.workspace_id, None)
            if os.getenv("DATABASE_URL"):
                from agent import storage
                ok = storage.delete_document(FILES, file_id)
                storage.delete_document(CHUNKS, file_id)
                storage.delete_blob(file_id)
                if not ok:
                    self._files[file_id] = kf
                    raise KnowledgeError("The file couldn't be deleted: the database is unavailable.", 503)
            else:
                self._write_local_index()
                for sub, ext in (("chunks", ".json"), ("blobs", ".bin")):
                    p = os.path.join(self.local_dir, sub, file_id + ext)
                    if os.path.isfile(p):
                        os.remove(p)
        return True

    # search -------------------------------------------------------------------------------------------
    def search(self, workspace_id: str, query: str, limit: int = 4) -> List[Dict[str, Any]]:
        """Best passages for a question, from this firm's ready files only."""
        self._refresh()
        with self._lock:
            idx = self._indexes.get(workspace_id)
            if idx is None:
                passages = []
                for fid, doc in self._chunks.items():
                    kf = self._files.get(fid)
                    if kf and kf.workspace_id == workspace_id and kf.status == "ready":
                        passages.extend(doc.get("chunks", []))
                idx = self._indexes[workspace_id] = _Index(passages)
        hits = idx.search(query, limit)
        top = hits[0][0] if hits else 1.0
        return [{"file_id": p["file_id"], "file_name": p["file_name"], "page": p.get("page") or None,
                 "text": p["text"], "score": round(s, 3), "relevance": round(s / top, 3)} for s, p in hits]


knowledge_manager = KnowledgeManager()
