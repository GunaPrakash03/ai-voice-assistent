#!/usr/bin/env python3
"""scripts/verify_knowledge.py — the per-firm knowledge base (web/knowledge-plan.html, Phases 1-2).

1. Offline (throwaway folder, DATABASE_URL cleared): file types are decided from the bytes (PDF, Word,
   Excel, CSV, Markdown, text, PNG/JPEG/WebP accepted; old .doc, zip, binary and empty refused); every
   format reads into the right text; a scanned PDF page and an image are read by a Gemini stand-in, and
   fail with a plain reason when no reader is set up; passages stay short, overlap and keep their page;
   search ranks the right file first and never crosses firms; limits, duplicates, delete, reprocess and a
   reload from disk.
2. PostgreSQL (only with VERIFY_PG_URL, a throwaway database): the same store round-trips through
   app_documents + knowledge_blobs, and a second manager (the worker) sees the files.
3. HTTP on an isolated server copy (own config, no .env, no database, no Gemini key): upload every
   format as a Product Admin, list, preview text, download the original byte-for-byte, search, delete,
   reprocess; Member Admins can read but not change; another firm's admin can't see or touch the files;
   a Super Admin can work on any firm with ?workspace_id=; oversize and wrong types are refused.
4. Calls (offline, Phase 2): each caller turn's instructions carry the firm's matching passages after the
   global rules (FIRM KNOWLEDGE), and none when nothing matches or for another firm; the global rule says to
   answer only from them; a follow-up keeps its topic; the knowledge tool searches the call's firm and
   says "call back" instead of guessing; the reply path (phone calls and Test Call) sends the passages to
   the model; the scripted fallback reads them; the worker refreshes instructions every turn.
   With VERIFY_LIVE_LLM=1 (uses the Gemini key in .env, 2 calls) a real model answers from an upload.
Never touches the real config, .env or database.
"""

import http.cookiejar
import io
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
passed = 0
failed = 0


def check(label, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  [PASS] {label}")
    else:
        failed += 1
        print(f"  [FAIL] {label} {detail}")


# ── Sample files, built here so the suite needs no fixtures ──────────────────────────────────────────
def make_pdf(scanned_page=True):
    import fitz
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 72), "Bottini Law fee schedule.", fontsize=12)
    p.insert_text((72, 96), "Initial consultations are free for personal injury matters.", fontsize=11)
    p.insert_text((72, 120), "Estate planning starts at 1,500 dollars for a simple will.", fontsize=11)
    if scanned_page:
        p2 = doc.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 80), 0)
        pix.clear_with(255)
        p2.insert_image(fitz.Rect(72, 72, 272, 152), pixmap=pix)   # an image, no text: a "scan"
    out = doc.tobytes()
    doc.close()
    return out


def make_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("Office hours: Monday to Friday, 8:30 am to 5:30 pm. Parking is free in the garage on Elm Street.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Attorney", "Practice"
    t.cell(1, 0).text, t.cell(1, 1).text = "Dana Reyes", "Workers compensation"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Locations"
    ws.append(["City", "Address", "Phone"])
    ws.append(["San Diego", "7817 Ivanhoe Avenue", "858-914-2001"])
    ws.append(["Riverside", "3801 University Ave", "951-555-0100"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_image(fmt="PNG"):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (320, 120), "white")
    ImageDraw.Draw(im).text((10, 50), "Spanish line: press 2", fill="black")
    buf = io.BytesIO()
    im.save(buf, fmt)
    return buf.getvalue()


def make_zip():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes.txt", "hello")
    return buf.getvalue()


SAMPLES = {
    "fees.pdf": (make_pdf, "pdf"),
    "office.docx": (make_docx, "docx"),
    "locations.xlsx": (make_xlsx, "xlsx"),
    "faq.csv": (lambda: b"question,answer\nDo you take Spanish calls?,Yes - press 2 for Spanish\n", "csv"),
    "intake.md": (lambda: "# Intake\n\nWe handle **dog bite** claims across California.\n".encode(), "markdown"),
    "notes.txt": (lambda: "﻿The firm was founded in 1984 by Francis Bottini.".encode("utf-8"), "text"),
    "flyer.png": (lambda: make_image("PNG"), "image"),
    "photo.jpg": (lambda: make_image("JPEG"), "image"),
}


def fake_reader(mime, data):
    return "Image: a sign.\n\nSpanish line: press 2. Scanned retainer terms: 33 percent contingency fee."


# ── 1. Offline ───────────────────────────────────────────────────────────────────────────────────────
def offline():
    print("\n1. Offline: types, reading, passages, search, limits")
    reqs = open(os.path.join(ROOT_DIR, "requirements.txt")).read().lower()
    for pkg in ("pymupdf", "python-docx", "openpyxl", "pillow"):
        check(f"requirements.txt installs {pkg} (deploys read files with it)", pkg in reqs)
    os.environ["DATABASE_URL"] = ""
    sys.path.insert(0, ROOT_DIR)
    import agent.knowledge_manager as km
    tmp = tempfile.mkdtemp(prefix="knowledge-")
    try:
        files = {n: f() for n, (f, _) in SAMPLES.items()}
        for name, (_, kind) in SAMPLES.items():
            check(f"sniff {name} -> {kind}", km.sniff(files[name], name)[0] == kind, km.sniff(files[name], name))
        check("type comes from the bytes: a PDF named .txt is a PDF", km.sniff(files["fees.pdf"], "x.txt")[0] == "pdf")
        check("WebP accepted", km.sniff(b"RIFF\x00\x00\x00\x00WEBPVP8 ", "a.webp")[0] == "image")
        for label, data, needle in (("old .doc refused", b"\xd0\xcf\x11\xe0" + b"\x00" * 100, ".docx"),
                                    ("zip refused", make_zip(), "zip"),
                                    ("binary refused", b"\x00\x01\x02binary", "isn't supported"),
                                    ("empty refused", b"", "empty"),
                                    ("non-UTF-8 text refused", "café".encode("latin-1"), "UTF-8")):
            try:
                km.sniff(data, "f.bin")
                check(label, False, "accepted")
            except km.KnowledgeError as e:
                check(label, needle.lower() in str(e).lower(), str(e))
        check("file name cleaned", km.clean_name("../../etc/\x07passwd") == "passwd" and km.clean_name("") == "Untitled")

        km.IMAGE_READER = fake_reader
        r = km.extract("pdf", "application/pdf", files["fees.pdf"])
        text = " ".join(t for t, _ in r["segments"])
        check("PDF text page read", "Initial consultations are free" in text, text[:120])
        check("PDF scanned page read by the image reader", "33 percent" in text and r["engine"] == "pdf-text+gemini", r["engine"])
        check("PDF pages kept", [p for _, p in r["segments"]] == [1, 2] and r["pages"] == 2)
        d = km.extract("docx", "", files["office.docx"])["segments"][0][0]
        check("Word paragraphs and table read", "Elm Street" in d and "Dana Reyes | Workers compensation" in d, d[:160])
        x = km.extract("xlsx", "", files["locations.xlsx"])["segments"][0][0]
        check("Excel sheet read as rows", "Sheet: Locations" in x and "Riverside | 3801 University Ave | 951-555-0100" in x, x)
        check("CSV read", "press 2 for Spanish" in km.extract("csv", "", files["faq.csv"])["segments"][0][0])
        check("text read without the BOM", km.extract("text", "", files["notes.txt"])["segments"][0][0].startswith("The firm"))
        check("image read by the image reader", "Spanish line" in km.extract("image", "image/png", files["flyer.png"])["segments"][0][0])

        km.IMAGE_READER = km._gemini_read_image
        os.environ.pop("GEMINI_API_KEY", None)
        os.environ.pop("GOOGLE_API_KEY", None)
        import agent.schema_extractor as se
        orig_key = se._gemini_key
        se._gemini_key = lambda: ""
        try:
            try:
                km.extract("image", "image/png", files["flyer.png"])
                check("image without a Gemini key fails", False, "accepted")
            except km.KnowledgeError as e:
                check("image without a Gemini key fails with the reason", "Gemini API key" in str(e), str(e))
            r = km.extract("pdf", "", files["fees.pdf"])
            check("PDF still reads its text pages without a key, and says a page was skipped",
                  "Initial consultations" in r["segments"][0][0] and "1 scanned page could not be read" in r["notes"], r["notes"])
            try:
                km.extract("pdf", "", _scan_only_pdf())
                check("scan-only PDF without a key fails", False, "accepted")
            except km.KnowledgeError as e:
                check("scan-only PDF without a key fails with the reason", "no Gemini API key" in str(e), str(e))
        finally:
            se._gemini_key = orig_key

        long = ("Sentence number %d talks about car accident claims in Los Angeles county. " * 60) % tuple(range(60))
        ps = km.split_passages([(long, 3), ("Short page.", 4)])
        check("passages are at most 900 characters", all(len(p["text"]) <= km.CHUNK_CHARS + km.CHUNK_OVERLAP for p in ps), max(len(p["text"]) for p in ps))
        check("passages overlap", len(ps) > 2 and ps[1]["text"].split()[0] in ps[0]["text"])
        check("passages keep their page and never cross pages", ps[-1] == {"i": len(ps) - 1, "text": "Short page.", "page": 4} and all(p["page"] == 3 for p in ps[:-1]))

        km.IMAGE_READER = fake_reader
        m = km.KnowledgeManager(local_dir=tmp)
        for name in SAMPLES:
            kf = m.upload("ws-a", name, files[name], uploaded_by="admin@a.test", wait=True)
            check(f"{name} ready", kf.status == "ready" and kf.chunk_count > 0, (kf.status, kf.error))
        other = m.upload("ws-b", "other.txt", b"Firm B secret: our retainer is 40 percent.", wait=True)
        check("usage counts the firm's files", m.usage("ws-a")["files"] == len(SAMPLES) and m.usage("ws-a")["ready"] == len(SAMPLES))

        def top(q, ws="ws-a"):
            hits = m.search(ws, q)
            return hits[0]["file_name"] if hits else None
        check("search: fees -> the PDF", top("How much does a will cost?") == "fees.pdf", m.search("ws-a", "How much does a will cost?")[:1])
        check("search: parking -> the Word file", top("Where can I park?") == "office.docx")
        check("search: Riverside phone -> the spreadsheet", top("phone number for the Riverside office") == "locations.xlsx")
        check("search: dog bite -> the Markdown file", top("do you handle dog bites") == "intake.md")
        check("search: founded -> the text file", top("when was the firm founded") == "notes.txt")
        check("search: opening hours -> the Word file", top("what time do you open?") == "office.docx", m.search("ws-a", "what time do you open?")[:2])
        check("search: located -> the spreadsheet", top("where are you located in San Diego") == "locations.xlsx")
        check("search: how much do you charge -> the PDF", top("how much do you charge for a will") == "fees.pdf")
        check("search: lawyer for workers comp -> the Word file", top("which lawyer does workers comp") == "office.docx")
        hit = m.search("ws-a", "consultation free injury")[0]
        check("search result carries file, page and passage", hit["file_name"] == "fees.pdf" and hit["page"] == 1 and "free" in hit["text"])
        check("search never crosses firms", not any(h["file_name"] == "other.txt" for h in m.search("ws-a", "retainer 40 percent secret")))
        check("the other firm finds its own file", top("retainer", ws="ws-b") == "other.txt")
        check("nothing for an unrelated question", m.search("ws-a", "zzzz qqqq") == [])

        for label, call, needle, status in (
                ("duplicate refused (409)", lambda: m.upload("ws-a", "copy.docx", files["office.docx"]), "already in the knowledge base", 409),
                ("oversize refused (413)", lambda: m.upload("ws-a", "big.txt", b"a" * (km.MAX_FILE_BYTES + 1)), "limit is 20 MB", 413),
                ("no firm refused", lambda: m.upload("", "x.txt", b"hello"), "No firm", 400)):
            try:
                call()
                check(label, False, "accepted")
            except km.KnowledgeError as e:
                check(label, needle in str(e) and e.status == status, (str(e), e.status))
        orig_firm = km.MAX_FIRM_BYTES
        km.MAX_FIRM_BYTES = m.usage("ws-a")["bytes"] + 10
        try:
            m.upload("ws-a", "more.txt", b"x" * 100)
            check("firm space limit refused", False, "accepted")
        except km.KnowledgeError as e:
            check("firm space limit refused", "Not enough space" in str(e) and e.status == 413, str(e))
        finally:
            km.MAX_FIRM_BYTES = orig_firm

        bg = m.upload("ws-a", "later.txt", b"Background read: the Fresno office opens in 2027.")
        for _ in range(50):
            if m.get(bg.file_id).status != "processing":
                break
            time.sleep(0.05)
        check("background read finishes", m.get(bg.file_id).status == "ready")

        km.IMAGE_READER = km._gemini_read_image
        se._gemini_key = lambda: ""
        try:
            img = m.upload("ws-a", "scan.png", make_image("PNG")[:-1] + b"\x00", wait=True)
            check("failed file recorded with its reason", img.status == "failed" and "Gemini" in img.error, (img.status, img.error))
            km.IMAGE_READER = fake_reader
            m.reprocess(img.file_id, wait=True)
            check("reprocess reads it once a reader is there", m.get(img.file_id).status == "ready")
        finally:
            se._gemini_key = orig_key

        check("original bytes kept", m.load_blob(m.list_files("ws-a")[-1].file_id) == files["fees.pdf"])
        pdf_id = next(k.file_id for k in m.list_files("ws-a") if k.name == "fees.pdf")
        check("preview text kept", "fee schedule" in m.text(pdf_id))
        m2 = km.KnowledgeManager(local_dir=tmp)
        check("reload from disk: same files and search", len(m2.list_files("ws-a")) == len(m.list_files("ws-a")) and (m2.search("ws-a", "park") or [{}])[0].get("file_name") == "office.docx")
        check("delete", m.delete(pdf_id) and m.get(pdf_id) is None and m.load_blob(pdf_id) is None and top("will cost") != "fees.pdf")
        check("delete unknown -> False", m.delete("kf-nope") is False)
        check("deleted file gone after reload", km.KnowledgeManager(local_dir=tmp).get(pdf_id) is None)
        check("other firm untouched", m.get(other.file_id).status == "ready")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _scan_only_pdf():
    import fitz
    doc = fitz.open()
    p = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 40), 0)
    pix.clear_with(255)
    p.insert_image(fitz.Rect(72, 72, 172, 112), pixmap=pix)
    out = doc.tobytes()
    doc.close()
    return out


# ── 4. Calls (Phase 2) ───────────────────────────────────────────────────────────────────────────────
def calls():
    print("\n4. Calls: knowledge on every caller turn")
    os.environ["DATABASE_URL"] = ""
    import asyncio
    import agent.knowledge_manager as km
    from agent import firm_context as fc
    tmp = tempfile.mkdtemp(prefix="knowledge-calls-")
    # The shared singleton is what the call paths use: point it at a throwaway folder.
    m = km.knowledge_manager
    orig_dir = m.local_dir
    m.local_dir = tmp
    m._load()
    try:
        m.upload("ws-default", "fees.pdf", make_pdf(scanned_page=False), wait=True)
        m.upload("ws-default", "office.docx", make_docx(), wait=True)
        m.upload("ws-other", "other.txt", b"Other firm: our retainer is 40 percent and parking is on Oak Street.", wait=True)

        block = fc.knowledge_context("ws-default", "how much is a simple will?")
        check("matching passage added", block.startswith(fc.KNOWLEDGE_MARK) and "1,500 dollars" in block and "(fees.pdf, page 1)" in block, block[:200])
        check("nothing added for an unrelated question", fc.knowledge_context("ws-default", "zzzz qqqq") == "")
        check("nothing added for an empty turn", fc.knowledge_context("ws-default", "") == "")
        check("another firm's files never added", "Oak Street" not in fc.knowledge_context("ws-default", "where do I park, Oak Street?"))
        check("agents with no firm use the default firm", "1,500" in fc.knowledge_context("", "price of a will"))
        orig_max = fc.KNOWLEDGE_MAX_CHARS
        fc.KNOWLEDGE_MAX_CHARS = 250
        check("block size capped", len(fc.knowledge_context("ws-default", "parking will fee hours")) < 250 + 300)
        fc.KNOWLEDGE_MAX_CHARS = orig_max

        ins = fc.call_instructions("You are Maya, the receptionist.", "ws-default", caller_text="where can I park?")
        check("order: global rules, agent prompt, then FIRM KNOWLEDGE",
              ins.index(fc.GLOBAL_MARK) < ins.index("You are Maya") < ins.index(fc.KNOWLEDGE_MARK + " (") and "garage on Elm Street" in ins)
        check("global rule: answer from FIRM KNOWLEDGE, never guess", "FIRM KNOWLEDGE sections below" in fc.GLOBAL_INSTRUCTIONS and "never guess a price" in fc.GLOBAL_INSTRUCTIONS)
        check("no caller text -> no knowledge block (call start)", fc.KNOWLEDGE_MARK + " (" not in fc.call_instructions("You are Maya.", "ws-default"))

        from agent.agent_builder import knowledge_query, agent_builder, AgentConfig
        hist = [{"speaker": "agent", "text": "Hi, this is Maya."}, {"speaker": "caller", "text": "I need a simple will."},
                {"speaker": "agent", "text": "Sure."}, {"speaker": "caller", "text": "And how much is that?"}]
        check("follow-up keeps its topic", knowledge_query("And how much is that?", hist) == "I need a simple will. And how much is that?")
        check("latest turn added when not yet in history", knowledge_query("Where do I park?", hist[:2]) == "I need a simple will. Where do I park?")
        check("follow-up finds the fee", "1,500" in fc.knowledge_context("ws-default", knowledge_query("And how much is that?", hist)))

        from agent.tool_manager import ToolRegistry, REAL_BUILTINS
        reg = ToolRegistry()
        tool = reg.get_tool("query_knowledge_base")
        check("a saved endpoint-less copy doesn't replace the real built-ins",
              all(reg.get_tool(n).handler.__name__ == n for n in REAL_BUILTINS), [reg.get_tool(n).handler.__name__ for n in REAL_BUILTINS])
        reg.workspace_id = "ws-default"
        res = asyncio.run(tool.handler(query="what are your office hours"))
        check("knowledge tool: real passage from the call's firm", res["found"] and "8:30 am" in res["snippet"] and res["sources"][0] == "office.docx", res)
        res = asyncio.run(tool.handler(query="zzzz qqqq"))
        check("knowledge tool: nothing found -> no guess", res["found"] is False and res["snippet"] == "")
        reg.workspace_id = "ws-other"
        res = asyncio.run(tool.handler(query="office hours parking"))
        check("knowledge tool: other firm gets only its own files", res["found"] and "office.docx" not in res["sources"], res.get("sources"))
        check("old made-up shop answers are gone", "30-day hassle-free" not in open(os.path.join(ROOT_DIR, "agent", "tool_manager.py")).read())
        import agent.llm_manager as lm
        fmt = lm.StreamingDialogueManager._format_grounded_tool_response
        check("tool reply when nothing found offers a call back",
              "call you back" in fmt(None, "query_knowledge_base", {"status": "success", "result": {"found": False}}))
        check("tool reply speaks the passage",
              "1,500 dollars" in fmt(None, "query_knowledge_base", {"status": "success", "result": {"found": True, "snippet": "A simple will is 1,500 dollars."}}))

        # The reply path used by phone calls (Twilio) and the Agent Builder Test Call: capture what is sent.
        import urllib.request as ur
        sent = {}

        class FakeResp:
            def __init__(self, body): self.body = body
            def read(self): return self.body
            def __enter__(self): return self
            def __exit__(self, *a): return False

        def fake_urlopen(req, timeout=None):
            sent["body"] = json.loads(req.data.decode())
            return FakeResp(json.dumps({"candidates": [{"content": {"parts": [{"text": "A simple will is fifteen hundred dollars."}]}}]}).encode())
        cfg = AgentConfig(agent_id="t", name="Maya", first_message="Thanks for calling.", system_prompt="You are Maya.", llm_model="gemini-3.5-flash-lite", workspace_id="ws-default")
        orig_open, orig_key = ur.urlopen, os.environ.get("GEMINI_API_KEY")
        ur.urlopen, os.environ["GEMINI_API_KEY"] = fake_urlopen, "test-key"
        try:
            reply = agent_builder._preview_reply(cfg, "How much does a simple will cost?", None, history=hist[:1] + [{"speaker": "caller", "text": "How much does a simple will cost?"}])
        finally:
            ur.urlopen = orig_open
            if orig_key is None:
                os.environ.pop("GEMINI_API_KEY", None)
            else:
                os.environ["GEMINI_API_KEY"] = orig_key
        system = sent.get("body", {}).get("system_instruction", {}).get("parts", [{}])[0].get("text", "")
        check("phone/Test Call reply: the model is sent the matching passage", fc.KNOWLEDGE_MARK in system and "1,500 dollars" in system, system[-300:])
        check("phone/Test Call reply returned", "fifteen hundred" in reply, reply)

        for k in ("GEMINI_API_KEY", "GOOGLE_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
            os.environ.pop(k, None)
        cfg2 = AgentConfig(agent_id="t2", name="Maya", first_message="Thanks for calling.", system_prompt="You are Maya.", llm_model="none", workspace_id="ws-default")
        r = agent_builder._preview_reply(cfg2, "what are your office hours", "query_knowledge_base")
        check("scripted fallback reads the passage", r.startswith("Here's what I have:") and "8:30 am" in r, r)
        r = agent_builder._preview_reply(cfg2, "zzzz qqqq", "query_knowledge_base")
        check("scripted fallback with nothing found offers a call back", "call you back" in r, r)

        src = open(os.path.join(ROOT_DIR, "agent", "worker.py")).read()
        turn = src[src.index("async def execute_llm_turn"):src.index("metrics = await llm_manager.generate_response(")]
        check("worker: every caller turn rebuilds the instructions with the knowledge query",
              "call_instructions, cfg_now.system_prompt, cfg_now.workspace_id, query" in turn and "asyncio.to_thread" in turn)
        check("worker: the knowledge tool searches the call's firm", "tool_registry.workspace_id = (active_cfg.workspace_id" in src)

        if os.getenv("VERIFY_LIVE_LLM") == "1":
            key = ""
            for line in open(os.path.join(ROOT_DIR, ".env")):
                if line.startswith("GEMINI_API_KEY="):
                    key = line.split("=", 1)[1].strip().strip("'\"")
            os.environ["GEMINI_API_KEY"] = key
            cfg3 = AgentConfig(agent_id="t3", name="Maya", first_message="Thanks for calling.", system_prompt="You are Maya, receptionist at Bottini Law.", llm_model="gemini-3.5-flash-lite", workspace_id="ws-default")
            r1 = agent_builder._preview_reply(cfg3, "Where can I park when I come in?", None,
                                              history=[{"speaker": "agent", "text": "Thanks for calling Bottini Law, this is Maya."},
                                                       {"speaker": "caller", "text": "Where can I park when I come in?"}])
            check("live model answers from the upload (parking)", "elm" in r1.lower() or "garage" in r1.lower(), r1)
            r2 = agent_builder._preview_reply(cfg3, "What does your firm charge for a divorce?", None,
                                              history=[{"speaker": "agent", "text": "Thanks for calling Bottini Law, this is Maya."},
                                                       {"speaker": "caller", "text": "What does your firm charge for a divorce?"}])
            check("live model doesn't invent a price it doesn't have", not any(c.isdigit() for c in r2), r2)
            os.environ.pop("GEMINI_API_KEY", None)
    finally:
        m.local_dir = orig_dir
        m._load()
        shutil.rmtree(tmp, ignore_errors=True)


# ── 2. PostgreSQL (throwaway database only) ──────────────────────────────────────────────────────────
def postgres():
    url = os.getenv("VERIFY_PG_URL", "")
    if not url:
        print("\n2. PostgreSQL: [SKIP] set VERIFY_PG_URL to a throwaway database to include these checks")
        return
    print("\n2. PostgreSQL (throwaway database)")
    os.environ["DATABASE_URL"] = url
    import agent.knowledge_manager as km
    from agent import storage
    storage._state.update(url=None, ok=False, checked=0.0, schema=False)
    km.IMAGE_READER = fake_reader
    m = km.KnowledgeManager(local_dir=tempfile.mkdtemp(prefix="knowledge-pg-"))
    docx_bytes = make_docx()
    kf = m.upload("ws-pg", "office.docx", docx_bytes, wait=True)
    check("upload stored and read", kf.status == "ready", (kf.status, kf.error))
    check("original bytes round-trip through knowledge_blobs", storage.load_blob(kf.file_id) == docx_bytes)
    worker = km.KnowledgeManager(local_dir=tempfile.mkdtemp(prefix="knowledge-pg2-"))
    check("a second process (the worker) sees the file and searches it",
          worker.get(kf.file_id) is not None and worker.search("ws-pg", "parking")[0]["file_name"] == "office.docx")
    m.delete(kf.file_id)
    check("delete removes record, passages and bytes",
          storage.load_document(km.FILES, kf.file_id) is None and storage.load_document(km.CHUNKS, kf.file_id) is None and storage.load_blob(kf.file_id) is None)
    os.environ["DATABASE_URL"] = ""


# ── 3. HTTP on an isolated server copy ───────────────────────────────────────────────────────────────
def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


class Client:
    def __init__(self, base):
        self.base = base
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def raw(self, method, path, data=None, headers=None):
        req = urllib.request.Request(self.base + path, data=data, method=method, headers=headers or {})
        try:
            with self.opener.open(req, timeout=60) as r:
                return r.status, r.read(), dict(r.headers)
        except urllib.error.HTTPError as e:
            return e.code, e.read(), dict(e.headers)

    def call(self, method, path, body=None):
        code, raw, _ = self.raw(method, path, json.dumps(body).encode() if body is not None else None, {"Content-Type": "application/json"})
        try:
            return code, json.loads(raw or b"{}")
        except ValueError:
            return code, {}

    def upload(self, name, data, ws="", wait=True):
        q = urllib.parse.urlencode({k: v for k, v in (("name", name), ("workspace_id", ws), ("wait", "1" if wait else "")) if v})
        code, raw, _ = self.raw("POST", "/api/v1/knowledge/upload?" + q, data, {"Content-Type": "application/octet-stream"})
        try:
            return code, json.loads(raw or b"{}")
        except ValueError:
            return code, {}

    def login(self, email, pw):
        code, body = self.call("POST", "/api/v1/auth/login", {"email": email, "password": pw})
        return code == 200 and body.get("step") == "done"


def http_suite():
    print("\n3. HTTP on an isolated server copy (no database, no Gemini key)")
    tmp = tempfile.mkdtemp(prefix="knowledge-http-")
    for d in ("agent", "scripts", "web"):
        shutil.copytree(os.path.join(ROOT_DIR, d), os.path.join(tmp, d), ignore=shutil.ignore_patterns("__pycache__", "audio", "*.mp3", "*.wav"))
    os.makedirs(os.path.join(tmp, "config"))
    os.makedirs(os.path.join(tmp, "recordings"))
    shutil.copy(os.path.join(ROOT_DIR, "config", "agents.json"), os.path.join(tmp, "config", "agents.json"))
    pw = "Passw0rd!x"
    env = {k: v for k, v in os.environ.items() if not k.startswith(("SMTP_", "RESEND_", "MAIL_", "TWILIO_", "TELNYX_", "GEMINI_", "GOOGLE_API"))}
    env.update(DATABASE_URL="", AUTH_TRUST_LOOPBACK="0", SUPER_ADMIN_EMAIL="root@platform.test", SUPER_ADMIN_PASSWORD=pw,
               LIVEKIT_URL="ws://127.0.0.1:1", LIVEKIT_API_KEY="d", LIVEKIT_API_SECRET="dummy-secret-dummy-secret")
    code = (
        "import sys; sys.path.insert(0, '.')\n"
        "from agent.auth_manager import AuthManager\n"
        "am = AuthManager(store_path='config/auth_store.json')\n"
        f"a = am.create_organization('Firm A', admin_email='admin@a.test', admin_password='{pw}')\n"
        f"b = am.create_organization('Firm B', admin_email='admin@b.test', admin_password='{pw}')\n"
        "ws = a['workspace_id'] if isinstance(a, dict) else a.workspace_id\n"
        f"u = am.create_user(ws, 'member@a.test', 'member_admin'); am.set_password(u.user_id, '{pw}')\n"
        "am._save_store(); print(ws)\n"
    )
    seeded = subprocess.run([sys.executable, "-c", code], cwd=tmp, env=env, capture_output=True, text=True)
    ws_a = (seeded.stdout.strip().splitlines() or [""])[-1]
    check("two firms and a Member Admin created", ws_a.startswith("ws-"), seeded.stderr[-300:])
    port = free_port()
    proc = subprocess.Popen([sys.executable, "scripts/serve.py", str(port)], cwd=tmp, env=env,
                            stdout=open(os.path.join(tmp, "server.log"), "w"), stderr=subprocess.STDOUT)
    base = f"http://127.0.0.1:{port}"
    try:
        for _ in range(80):
            try:
                urllib.request.urlopen(base + "/login", timeout=1)
                break
            except Exception:
                time.sleep(0.25)
        A, B, M, S = Client(base), Client(base), Client(base), Client(base)
        check("everyone signs in", A.login("admin@a.test", pw) and B.login("admin@b.test", pw) and M.login("member@a.test", pw) and S.login("root@platform.test", pw))
        anon = Client(base)
        code, _ = anon.call("GET", "/api/v1/knowledge")
        check("signed out: refused", code == 401, code)

        files = {n: f() for n, (f, _) in SAMPLES.items()}
        ids = {}
        for name, data in files.items():
            code, body = A.upload(name, data)
            f = body.get("file") or {}
            ids[name] = f.get("file_id")
            want = "failed" if SAMPLES[name][1] == "image" else "ready"
            check(f"upload {name} -> {want}", code == 201 and f.get("status") == want, (code, f.get("status"), f.get("error") or body.get("error")))
        code, body = A.call("GET", "/api/v1/knowledge")
        check("images fail with the Gemini reason (no key in this copy)",
              all("Gemini API key" in f["error"] for f in body["files"] if f["kind"] == "image"))
        check("PDF reads its text page and notes the scanned one",
              any(f["name"] == "fees.pdf" and "1 scanned page could not be read" in f["notes"] for f in body["files"]))
        check("list: the firm's files, usage, can_edit", len(body["files"]) == len(files) and body["usage"]["files"] == len(files) and body["can_edit"] is True)

        code, body = A.call("GET", "/api/v1/knowledge/search?q=" + urllib.parse.quote("where can I park"))
        check("search finds the Word file", code == 200 and body["results"] and body["results"][0]["file_name"] == "office.docx", body.get("results", [])[:1])
        code, body = A.call("GET", "/api/v1/knowledge/search?q=")
        check("search needs a question", code == 400)
        code, body = A.call("GET", "/api/v1/knowledge/text?file_id=" + ids["locations.xlsx"])
        check("preview text", code == 200 and "Riverside" in body["text"] and body["passages"] >= 1)
        code, raw, hdrs = A.raw("GET", "/api/v1/knowledge/download?file_id=" + ids["fees.pdf"])
        check("download returns the original bytes", code == 200 and raw == files["fees.pdf"], (code, len(raw)))
        check("download headers: attachment, type, nosniff",
              "attachment" in hdrs.get("Content-Disposition", "") and hdrs.get("Content-Type") == "application/pdf" and hdrs.get("X-Content-Type-Options") == "nosniff")

        for label, data, name, want in (("old .doc refused", b"\xd0\xcf\x11\xe0" + b"\x00" * 64, "a.doc", 400),
                                        ("zip refused", make_zip(), "a.zip", 400),
                                        ("empty refused", b"", "a.txt", 400),
                                        ("duplicate refused", files["office.docx"], "again.docx", 409)):
            code, body = A.upload(name, data)
            check(f"{label} ({want})", code == want and body.get("error"), (code, body))
        code, body = A.upload("big.txt", b"a" * (20 * 1024 * 1024 + 1))
        check("oversize refused (413) with the limit", code == 413 and "20 MB" in body.get("error", ""), (code, body))

        code, body = M.call("GET", "/api/v1/knowledge")
        check("Member Admin can list", code == 200 and len(body["files"]) == len(files) and body["can_edit"] is False)
        code, _ = M.call("GET", "/api/v1/knowledge/search?q=parking")
        check("Member Admin can search", code == 200)
        code, body = M.upload("m.txt", b"member upload")
        check("Member Admin can't upload (403)", code == 403, (code, body))
        code, _ = M.call("POST", "/api/v1/knowledge/delete", {"file_id": ids["notes.txt"]})
        check("Member Admin can't delete (403)", code == 403, code)

        code, body = B.call("GET", "/api/v1/knowledge")
        check("other firm sees an empty library", code == 200 and body["files"] == [])
        for label, call in (("text", lambda: B.call("GET", "/api/v1/knowledge/text?file_id=" + ids["notes.txt"])),
                            ("download", lambda: B.call("GET", "/api/v1/knowledge/download?file_id=" + ids["notes.txt"])),
                            ("delete", lambda: B.call("POST", "/api/v1/knowledge/delete", {"file_id": ids["notes.txt"]}))):
            code, _ = call()
            check(f"other firm's admin can't {label} it (404)", code == 404, code)
        code, body = B.call("GET", "/api/v1/knowledge/search?q=founded")
        check("other firm's search doesn't reach it", code == 200 and body["results"] == [])
        code, _ = B.call("GET", "/api/v1/knowledge?workspace_id=" + ws_a)
        check("Product Admin can't pick another firm (403)", code == 403, code)
        code, body = B.upload("x.txt", b"cross-firm upload", ws=ws_a)
        check("Product Admin can't upload into another firm (403)", code == 403, code)

        code, body = S.call("GET", "/api/v1/knowledge?workspace_id=" + ws_a)
        check("Super Admin lists any firm", code == 200 and len(body["files"]) == len(files))
        code, body = S.upload("from-super.txt", b"Added by the platform team: holiday hours are 9 to 1.", ws=ws_a)
        check("Super Admin uploads into a firm", code == 201 and body["file"]["workspace_id"] == ws_a)
        code, _ = S.call("GET", "/api/v1/knowledge?workspace_id=ws-nope")
        check("unknown firm (404)", code == 404)

        code, body = A.call("POST", "/api/v1/knowledge/reprocess", {"file_id": ids["notes.txt"], "wait": True})
        check("reprocess", code == 200 and body["file"]["status"] == "ready", (code, body))
        code, body = A.call("POST", "/api/v1/knowledge/delete", {"file_id": ids["notes.txt"]})
        check("delete", code == 200 and body.get("deleted") is True)
        code, _ = A.call("GET", "/api/v1/knowledge/text?file_id=" + ids["notes.txt"])
        check("deleted file is gone (404)", code == 404)

        proc.terminate()
        proc.wait(timeout=10)
        proc = subprocess.Popen([sys.executable, "scripts/serve.py", str(port)], cwd=tmp, env=env,
                                stdout=open(os.path.join(tmp, "server2.log"), "w"), stderr=subprocess.STDOUT)
        for _ in range(80):
            try:
                urllib.request.urlopen(base + "/login", timeout=1)
                break
            except Exception:
                time.sleep(0.25)
        A2 = Client(base)
        A2.login("admin@a.test", pw)
        code, body = A2.call("GET", "/api/v1/knowledge")
        check("after a restart: files and search still there",
              code == 200 and len(body["files"]) == len(files) and A2.call("GET", "/api/v1/knowledge/search?q=park")[1]["results"][0]["file_name"] == "office.docx")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()
        if failed:
            print("  server log:", os.path.join(tmp, "server.log"))
        else:
            shutil.rmtree(tmp, ignore_errors=True)


def main():
    offline()
    calls()
    postgres()
    if "--offline" not in sys.argv:
        http_suite()
    print(f"\n{passed}/{passed + failed} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
